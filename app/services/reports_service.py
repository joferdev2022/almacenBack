"""Reportes de solo lectura: ventas por fecha de venta, nunca por saldo/cobros."""
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from math import ceil
from typing import Literal
from unicodedata import combining, normalize

from bson.decimal128 import Decimal128
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.db.mongo import productsDb, salesDb
from app.utils.cash_helpers import LIMA, date_filter

CENT = Decimal("0.01")
MAX_LINES = 100_000
MAX_EXCEL_LINES = 25_000
NO_CATEGORY = "__sin_categoria__"
MONTH_NAMES = ("enero", "febrero", "marzo", "abril", "mayo", "junio",
               "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre")


@dataclass(frozen=True)
class ReportFilters:
    fecha_desde: date
    fecha_hasta: date
    producto: str = ""
    categoria: str = ""
    ventas: Literal["todas", "pagadas"] = "todas"

    def __post_init__(self):
        if self.ventas not in ("todas", "pagadas"):
            raise HTTPException(422, "El filtro de ventas debe ser todas o pagadas.")
        date_filter(self.fecha_desde, self.fecha_hasta)
        months = (self.fecha_hasta.year - self.fecha_desde.year) * 12 + self.fecha_hasta.month - self.fecha_desde.month + 1
        if months > 60:
            raise HTTPException(422, "Consulta como máximo 60 meses por reporte.")
        if self.fecha_desde.year < 1900:
            raise HTTPException(422, "La fecha desde debe ser posterior a 1899.")

    def serialize(self):
        return {"fecha_desde": self.fecha_desde.isoformat(), "fecha_hasta": self.fecha_hasta.isoformat(),
                "producto": self.producto, "categoria": self.categoria, "ventas": self.ventas}


def _search(value):
    return "".join(char for char in normalize("NFKD", str(value or "")).casefold() if not combining(char))


def _number(value):
    try:
        result = value.to_decimal() if isinstance(value, Decimal128) else Decimal(str(value))
        return result if result.is_finite() else None
    except (InvalidOperation, TypeError, ValueError):
        return None


def _amount(value, precision=CENT):
    return float(value.quantize(precision, rounding=ROUND_HALF_UP))


def _empty_sales_notice(count):
    return (f"Ventas vacías omitidas: {count}. Sin productos, con total y total original de S/ 0.00. "
            "Recuento por local, fecha y tipo de ventas, sin asignación de producto o categoría.")


def _months(filters):
    year, month = filters.fecha_desde.year, filters.fecha_desde.month
    rows = {}
    while (year, month) <= (filters.fecha_hasta.year, filters.fecha_hasta.month):
        key = f"{year:04d}-{month:02d}"
        rows[key] = {"mes": key, "cantidad": Decimal(0), "importeVendido": Decimal(0), "ventas": set(),
                     "ventasVaciasOmitidas": 0}
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    return rows


def build_report(user, filters: ReportFilters, for_export=False):
    # No matching por nombre: una importación puede reutilizar nombres con IDs distintos.
    catalog = {str(product["_id"]): product for product in productsDb.find(
        {"local": user["local"]}, {"nombre": 1, "categoria": 1, "precioCompra": 1,
                                 "unidadDeMedida": 1, "marca": 1, "descripcion": 1})}
    categories = sorted({str(product["categoria"]) for product in catalog.values() if product.get("categoria")}, key=_search)
    query = {"local": user["local"], "anulado": {"$ne": True},
             "fechaVenta": date_filter(filters.fecha_desde, filters.fecha_hasta)}
    projection = {"fechaVenta": 1, "estado": 1, "productos": 1, "precioTotal": 1, "precioTotalOriginal": 1}
    if filters.ventas == "pagadas":
        # "cancelado" significa totalmente pagada; anulado=true se excluye arriba.
        query["estado"] = "cancelado"
    months, groups, top, sales = _months(filters), {}, {}, set()
    quantity_total, amount_total = Decimal(0), Decimal(0)
    missing_cost, line_count, empty_sales = 0, 0, 0
    max_lines = min(MAX_LINES, MAX_EXCEL_LINES) if for_export else MAX_LINES
    search = _search(filters.producto.strip())
    for sale in salesDb.find(query, projection).sort([("fechaVenta", 1), ("_id", 1)]):
        sale_id = str(sale["_id"])
        instant = sale["fechaVenta"]
        instant = (instant if instant.tzinfo else instant.replace(tzinfo=timezone.utc)).astimezone(LIMA)
        month = instant.strftime("%Y-%m")
        items = sale.get("productos")
        # Solo una lista explícitamente vacía y ambos importes numéricos exactamente cero.
        # No inferir cero de campos ausentes ni de un saldo cobrado: el original también debe ser cero.
        if isinstance(items, list) and not items and all(
                _number(sale.get(field)) == 0 for field in ("precioTotal", "precioTotalOriginal")):
            empty_sales += 1
            months[month]["ventasVaciasOmitidas"] += 1
            continue
        if not isinstance(items, list) or not items:
            raise HTTPException(422, f"La venta {sale_id} no tiene un detalle de productos válido. Revísala antes de generar el reporte.")
        for item in items:
            if not isinstance(item, dict):
                raise HTTPException(422, f"La venta {sale_id} contiene un detalle inválido.")
            product_id = str(item.get("productoId") or "")
            product = catalog.get(product_id)
            name = str(item.get("productName") or (product or {}).get("nombre") or "Producto sin nombre")
            category = str(product["categoria"]) if product and product.get("categoria") else None
            if search and search not in _search(" ".join((product_id, name, str((product or {}).get("nombre", ""))))):
                continue
            if filters.categoria and filters.categoria != (category or NO_CATEGORY):
                continue
            quantity, price = _number(item.get("cantidad")), _number(item.get("precioUnitario"))
            if quantity is None or quantity <= 0 or price is None or price < 0:
                raise HTTPException(422, f"La venta {sale_id} tiene cantidad o precio inválido en {name}. Corrígela antes de generar el reporte.")
            line_count += 1
            if line_count > max_lines:
                limit = f"{max_lines:,}".replace(",", " ")
                raise HTTPException(422, f"El {'Excel' if for_export else 'reporte'} supera {limit} líneas. Reduce el período o filtra un producto.")
            amount = (quantity * price).quantize(CENT, rounding=ROUND_HALF_UP)
            cost = _number(product.get("precioCompra")) if product else None
            if cost is not None and cost < 0:
                cost = None
            missing_cost += cost is None
            base = {"productoId": product_id, "producto": name, "categoriaActual": category,
                    "presentacion": str((product or {}).get("unidadDeMedida") or ""),
                    "marca": str((product or {}).get("marca") or ""),
                    "descripcion": str((product or {}).get("descripcion") or ""),
                    "precioCompraActual": float(cost) if cost is not None else None}
            # Para documentos antiguos sin ID, el nombre agrupa solo ese histórico.
            key = product_id or "sin-id:" + name
            row = groups.setdefault((month, key), {**base, "mes": month, "cantidad": Decimal(0), "importeVendido": Decimal(0)})
            ranking = top.setdefault(key, {**base, "cantidad": Decimal(0), "importeVendido": Decimal(0)})
            for target in (row, ranking, months[month]):
                target["cantidad"] += quantity
                target["importeVendido"] += amount
            months[month]["ventas"].add(sale_id)
            sales.add(sale_id)
            quantity_total += quantity
            amount_total += amount

    rows = [{**row, "cantidad": float(row["cantidad"]), "importeVendido": _amount(row["importeVendido"]),
             "precioVentaPromedio": _amount(row["importeVendido"] / row["cantidad"], Decimal("0.0001"))}
            for row in sorted(groups.values(), key=lambda r: (r["mes"], _search(r["producto"]), r["productoId"]))]
    rankings = [{**row, "cantidad": float(row["cantidad"]), "importeVendido": _amount(row["importeVendido"])} for row in top.values()]
    monthly = [{"mes": row["mes"], "cantidad": float(row["cantidad"]), "importeVendido": _amount(row["importeVendido"]),
                "numeroVentas": len(row["ventas"]), "ventasVaciasOmitidas": row["ventasVaciasOmitidas"]}
               for row in months.values()]
    scope = "Solo ventas completamente pagadas" if filters.ventas == "pagadas" else "Todas las ventas, incluyendo crédito"
    notice = f"{scope}, por fecha de venta (America/Lima). Se excluyen anuladas. Los importes no representan cobros de caja."
    notices = [notice]
    if empty_sales:
        notices.append(_empty_sales_notice(empty_sales))
    return {"filtros": filters.serialize(), "local": user["local"], "zonaHoraria": "America/Lima",
            "generadoEn": datetime.now(LIMA).isoformat(), "avisos": notices, "categorias": categories,
            "resumen": {"importeVendido": _amount(amount_total), "unidadesVendidas": float(quantity_total),
                        "numeroVentas": len(sales), "lineasSinPrecioCompra": missing_cost,
                        "ventasVaciasOmitidas": empty_sales},
            "mensual": monthly, "filas": rows, "totalFilas": len(rows),
            "topCantidad": sorted(rankings, key=lambda r: (-r["cantidad"], _search(r["producto"]), r["productoId"]))[:10],
            "topImporte": sorted(rankings, key=lambda r: (-r["importeVendido"], _search(r["producto"]), r["productoId"]))[:10]}


def _write_cell(sheet, row, column, value):
    cell = sheet.cell(row=row, column=column)
    if isinstance(value, str):
        # Nombres/IDs son texto, nunca fórmulas ni vínculos ejecutables.
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell.data_type = "s"
    else:
        cell.value = value
    return cell


def export_excel(report):
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Ventas mensuales por producto"
    filters = report["filtros"]
    sale_scope = "Solo pagadas" if filters["ventas"] == "pagadas" else "Todas las ventas (incluye crédito)"
    context = (f"Local: {report['local']} | Rango aplicado: {filters['fecha_desde']} a {filters['fecha_hasta']} | "
               f"Ventas: {sale_scope} | Producto: {filters['producto'] or 'Todos'} | Categoría: "
               f"{'Sin categoría disponible' if filters['categoria'] == NO_CATEGORY else filters['categoria'] or 'Todas'}")
    headers = ["CANTIDAD", "PRODUCTO", "PRESENTACION", "MARCA", "CATEGORIA",
               "DESCRIPCION", "P. COMPRA", "P. VENTA", "IMPORTE VENDIDO"]
    widths = [14, 36, 24, 22, 24, 44, 20, 20, 24]
    formats = {1: "#,##0.###", 7: "#,##0.00######", 8: "#,##0.00##", 9: "#,##0.00"}
    rows_by_month = {month["mes"]: [] for month in report["mensual"]}
    for row in report["filas"]:
        rows_by_month[row["mes"]].append(row)

    for month in report["mensual"]:
        year, month_number = month["mes"].split("-")
        sheet_name = MONTH_NAMES[int(month_number) - 1] + year
        sheet = workbook.create_sheet(sheet_name)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        for column, header in enumerate(headers, 1):
            cell = _write_cell(sheet, 1, column, header)
            cell.fill = PatternFill("solid", fgColor="166534")
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(column)].width = widths[column - 1]
        sheet.row_dimensions[1].height = 32
        sheet["H1"].comment = Comment(
            "Precio de venta promedio ponderado: importe vendido dividido entre cantidad del producto en este mes.",
            "Reportes",
        )
        rows = rows_by_month[month["mes"]]
        for row_number, row in enumerate(rows, 2):
            values = [row["cantidad"], row["producto"], row["presentacion"], row["marca"],
                      row["categoriaActual"] or "", row["descripcion"],
                      row["precioCompraActual"] if row["precioCompraActual"] is not None else "No disponible",
                      row["precioVentaPromedio"], row["importeVendido"]]
            lines = 1
            for column, value in enumerate(values, 1):
                cell = _write_cell(sheet, row_number, column, value)
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="center", wrap_text=isinstance(value, str))
                if column in formats:
                    cell.number_format = formats[column]
                if isinstance(value, str):
                    lines = max(lines, ceil(len(value) / (widths[column - 1] - 2)))
            sheet.row_dimensions[row_number].height = max(28, min(120, lines * 15 + 6))
        if rows:
            table = Table(displayName=f"Ventas_{year}_{month_number}", ref=f"A1:I{len(rows) + 1}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
            sheet.add_table(table)
        else:
            sheet.merge_cells("B2:I2")
            _write_cell(sheet, 2, 2, "Sin ventas para los filtros seleccionados.")
            sheet.auto_filter.ref = "A1:I1"
        total_row = max(4, len(rows) + 3)
        for column, value in ((1, month["cantidad"]), (2, "TOTAL DEL MES"), (9, month["importeVendido"])):
            cell = _write_cell(sheet, total_row, column, value)
            cell.font = Font(name="Calibri", bold=True, color="166534")
            cell.fill = PatternFill("solid", fgColor="E8F5EC")
            cell.number_format = formats.get(column, "General")
        # Los filtros quedan fuera de la tabla; la primera fila siempre tiene las columnas solicitadas.
        footer = [(total_row + 2, context),
                  (total_row + 3, f"Generado: {report['generadoEn']} | Hora de Lima | Importes en soles (S/)")]
        if month["ventasVaciasOmitidas"]:
            footer.append((total_row + 4, _empty_sales_notice(month["ventasVaciasOmitidas"])))
        for row_number, text in footer:
            sheet.merge_cells(f"A{row_number}:I{row_number}")
            cell = _write_cell(sheet, row_number, 1, text)
            cell.font = Font(name="Calibri", size=10, color="64748B")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row_number].height = 28 if row_number == total_row + 3 else 45
        sheet.print_title_rows = "1:1"
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f"A1:I{footer[-1][0]}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()
