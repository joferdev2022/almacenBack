"""API y Excel de Reportes con datos aislados; nunca conecta a MongoDB."""
import json
import unittest
from copy import deepcopy
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from unittest.mock import patch
from urllib.parse import urlencode

from bson import ObjectId
from bson.decimal128 import Decimal128
from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook

from memory_db import MemoryDatabase, database_patches
from app.routes.reports_route import router
from app.services import reports_service as service
from app.services.auth_service import create_access_token


class ReportsTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.db = MemoryDatabase()
        for item in database_patches(self.db) + [patch.object(service, "salesDb", self.db.salesDb),
                                                patch.object(service, "productsDb", self.db.productsDb)]:
            item.start()
            self.addCleanup(item.stop)
        self.user = {"_id": ObjectId(), "username": "reports-admin", "permissions": 1, "local": 1}
        self.reader = {"_id": ObjectId(), "username": "reports-reader", "permissions": 0, "local": 1}
        self.foreign = {"_id": ObjectId(), "username": "reports-foreign", "permissions": 1, "local": 2}
        for user in (self.user, self.reader, self.foreign):
            self.db.authDb.insert_one(user)
        self.product = {"_id": ObjectId(), "local": 1, "nombre": "Café actual", "categoria": "Alimentos", "precioCompra": 15}
        self.other = {"_id": ObjectId(), "local": 1, "nombre": "Bolsa", "categoria": "Accesorios", "precioCompra": 0}
        self.db.productsDb.insert_one(self.product)
        self.db.productsDb.insert_one(self.other)
        self.filters = service.ReportFilters(date(2026, 1, 1), date(2026, 2, 28))
        self.app = FastAPI()
        self.app.include_router(router, prefix="/api")

    def line(self, quantity=2, price=20, product=None, name="Café histórico", **overrides):
        return {"productoId": str((product or self.product)["_id"]), "productName": name,
                "cantidad": quantity, "precioUnitario": price, "precioBuy": 14, **overrides}

    def sale(self, items=None, instant="2026-01-15T15:00:00+00:00", **overrides):
        row = {"_id": ObjectId(), "local": 1, "fechaVenta": datetime.fromisoformat(instant),
               "estado": "cancelado", "precioTotal": 999, "productos": items if items is not None else [self.line()], **overrides}
        self.db.salesDb.insert_one(row)
        return row

    def seed(self):
        self.sale([self.line(), self.line(1, 6, self.other, "Bolsa")])
        self.sale([self.line(3, 22)], instant="2026-02-01T04:59:59+00:00", estado="credito", precioTotal=1)
        self.sale([self.line(100, 999)], anulado=True)
        self.sale([self.line(100, 999)], local=2)

    async def request(self, path="/api/reports/sales", user="default", token=None, **params):
        identity = self.user if user == "default" else user
        headers = []
        if identity is not None:
            access = token or create_access_token({"sub": identity["username"], "local": identity["local"]}, timedelta(minutes=5))
            headers.append((b"authorization", ("Bearer " + access).encode()))
        messages = []
        async def receive():
            return {"type": "http.request", "body": b"", "more_body": False}
        async def send(message):
            messages.append(message)
        scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": "GET",
                 "scheme": "http", "path": path, "raw_path": path.encode(), "root_path": "",
                 "query_string": urlencode({"fecha_desde": "2026-01-01", "fecha_hasta": "2026-02-28", **params}).encode(),
                 "headers": headers, "client": ("127.0.0.1", 12345), "server": ("test", 80)}
        await self.app(scope, receive, send)
        start = next(message for message in messages if message["type"] == "http.response.start")
        body = b"".join(message.get("body", b"") for message in messages if message["type"] == "http.response.body")
        return start["status"], dict(start["headers"]), body

    def test_totals_use_lines_and_include_credit_not_balance(self):
        self.seed()
        report = service.build_report(self.user, self.filters)
        self.assertEqual(report["resumen"], {"importeVendido": 112, "unidadesVendidas": 6,
                                           "numeroVentas": 2, "lineasSinPrecioCompra": 0, "ventasVaciasOmitidas": 0})
        self.assertEqual(report["mensual"][0]["importeVendido"], 112)
        self.assertEqual(report["mensual"][0]["numeroVentas"], 2)
        self.assertEqual(report["mensual"][1], {"mes": "2026-02", "cantidad": 0, "importeVendido": 0, "numeroVentas": 0,
                                               "ventasVaciasOmitidas": 0})
        coffee = next(row for row in report["filas"] if row["productoId"] == str(self.product["_id"]))
        self.assertEqual(coffee["precioVentaPromedio"], 21.2)
        self.assertEqual(coffee["precioCompraActual"], 15)
        self.assertEqual(coffee["producto"], "Café histórico")

    def test_date_range_inclusive_in_lima_and_naive_utc(self):
        self.sale(instant="2026-01-01T04:59:59+00:00")
        self.sale(instant="2026-01-01T05:00:00+00:00")
        self.sale(instant="2026-02-01T04:59:59")
        self.sale(instant="2026-02-01T05:00:00+00:00")
        report = service.build_report(self.user, service.ReportFilters(date(2026, 1, 1), date(2026, 1, 31)), True)
        self.assertEqual(report["resumen"]["numeroVentas"], 2)
        self.assertEqual(report["mensual"], [{"mes": "2026-01", "cantidad": 4, "importeVendido": 80, "numeroVentas": 2,
                                             "ventasVaciasOmitidas": 0}])
        self.assertEqual(report["filas"][0]["mes"], "2026-01")

    def test_product_search_filters_lines_not_entire_basket(self):
        self.seed()
        for search in ("cafe", "ACTUAL", str(self.product["_id"])):
            report = service.build_report(self.user, service.ReportFilters(date(2026, 1, 1), date(2026, 2, 28), search))
            self.assertEqual(report["resumen"]["importeVendido"], 106)
            self.assertEqual(report["resumen"]["unidadesVendidas"], 5)
            self.assertEqual(report["resumen"]["numeroVentas"], 2)

    def test_category_filter_uses_current_catalog(self):
        self.seed()
        report = service.build_report(self.user, service.ReportFilters(date(2026, 1, 1), date(2026, 2, 28), categoria="Accesorios"))
        self.assertEqual(report["resumen"]["importeVendido"], 6)
        self.assertEqual(report["filas"][0]["precioCompraActual"], 0)

    def test_deleted_product_remains_and_never_matches_same_name(self):
        deleted_id = ObjectId()
        self.sale([self.line(product={"_id": deleted_id}, name="Café actual")])
        report = service.build_report(self.user, self.filters)
        self.assertEqual(report["filas"][0]["productoId"], str(deleted_id))
        self.assertIsNone(report["filas"][0]["precioCompraActual"])
        self.assertIsNone(report["filas"][0]["categoriaActual"])
        self.assertEqual(report["resumen"]["lineasSinPrecioCompra"], 1)
        unknown = service.build_report(self.user, service.ReportFilters(date(2026, 1, 1), date(2026, 2, 28), categoria=service.NO_CATEGORY))
        self.assertEqual(unknown["resumen"]["importeVendido"], 40)

    def test_foreign_product_does_not_disclose_cost(self):
        foreign_product = {"_id": ObjectId(), "local": 2, "nombre": "Otro", "precioCompra": 900, "categoria": "Privada"}
        self.db.productsDb.insert_one(foreign_product)
        self.sale([self.line(product=foreign_product)])
        report = service.build_report(self.user, self.filters)
        self.assertIsNone(report["filas"][0]["precioCompraActual"])
        self.assertNotIn("Privada", report["categorias"])

    def test_same_name_different_ids_stay_distinct(self):
        self.sale([self.line(), self.line(product=self.other)])
        self.assertEqual(len(service.build_report(self.user, self.filters)["filas"]), 2)

    def test_decimal_arithmetic_and_distinct_sale_count(self):
        self.sale([self.line(3, Decimal128("0.10")), self.line(2, "0.20")])
        report = service.build_report(self.user, self.filters)
        self.assertEqual(report["resumen"]["importeVendido"], 0.7)
        self.assertEqual(report["resumen"]["numeroVentas"], 1)
        self.assertEqual(report["filas"][0]["precioVentaPromedio"], 0.14)

    def test_invalid_legacy_detail_fails_clearly_instead_of_silent_totals(self):
        for items in ([], [self.line(price="NaN")], [self.line(quantity=-1)], ["broken"]):
            self.db.salesDb.documents.clear()
            self.sale(items)
            with self.assertRaises(HTTPException) as failure:
                service.build_report(self.user, self.filters)
            self.assertEqual(failure.exception.status_code, 422)

    def test_empty_report_has_zero_months_and_totals(self):
        report = service.build_report(self.user, self.filters)
        self.assertEqual(report["resumen"]["numeroVentas"], 0)
        self.assertEqual(len(report["mensual"]), 2)
        self.assertEqual(report["filas"], [])
        self.assertEqual(report["topCantidad"], [])
        self.assertEqual(report["resumen"]["ventasVaciasOmitidas"], 0)
        self.assertFalse(any("omitidas" in notice for notice in report["avisos"]))

    def test_four_empty_zero_sales_do_not_block_or_change_valid_results(self):
        self.seed()
        baseline = service.build_report(self.user, self.filters)
        for value in (0, 0.0, "0.00", Decimal128("0")):
            self.sale([], precioTotal=value, precioTotalOriginal=value)
        before = deepcopy((self.db.salesDb.documents, self.db.productsDb.documents))
        with patch.object(self.db.salesDb, "find", wraps=self.db.salesDb.find) as find:
            report = service.build_report(self.user, self.filters)
            projection = find.call_args.args[1]
            self.assertEqual(projection["precioTotal"], 1)
            self.assertEqual(projection["precioTotalOriginal"], 1)
        self.assertEqual(report["resumen"]["ventasVaciasOmitidas"], 4)
        self.assertEqual(report["mensual"][0]["ventasVaciasOmitidas"], 4)
        for key in ("importeVendido", "unidadesVendidas", "numeroVentas", "lineasSinPrecioCompra"):
            self.assertEqual(report["resumen"][key], baseline["resumen"][key])
        for key in ("filas", "totalFilas", "topCantidad", "topImporte"):
            self.assertEqual(report[key], baseline[key])
        self.assertIn("Ventas vacías omitidas: 4.", report["avisos"][1])
        self.assertIn("sin asignación de producto o categoría", report["avisos"][1])
        exported = service.build_report(self.user, self.filters, for_export=True)
        self.assertEqual(exported["resumen"], report["resumen"])
        service.export_excel(exported)
        self.assertEqual(before, (self.db.salesDb.documents, self.db.productsDb.documents))

    def test_empty_sales_with_nonzero_or_ambiguous_totals_still_fail(self):
        for field in ("precioTotal", "precioTotalOriginal"):
            for value in (1, -1, "0.001", "-0.001", None, "", "NaN", float("nan"), "Infinity", False, True, {}, []):
                with self.subTest(field=field, value=value):
                    self.db.salesDb.documents.clear()
                    sale = self.sale([], precioTotal=0, precioTotalOriginal=0)
                    self.db.salesDb.documents[0][field] = value
                    for for_export in (False, True):
                        with self.assertRaises(HTTPException) as error:
                            service.build_report(self.user, self.filters, for_export=for_export)
                        self.assertEqual(error.exception.status_code, 422)
                        self.assertIn(str(sale["_id"]), error.exception.detail)
            self.db.salesDb.documents.clear()
            self.sale([], precioTotal=0, precioTotalOriginal=0)
            del self.db.salesDb.documents[0][field]
            with self.assertRaises(HTTPException):
                service.build_report(self.user, self.filters)

    def test_missing_or_malformed_products_are_not_treated_as_empty_zero_sales(self):
        for products in (None, {}, "", "[]", False, ["broken"], [{}]):
            with self.subTest(products=products):
                self.db.salesDb.documents.clear()
                self.sale(productos=products, precioTotal=0, precioTotalOriginal=0)
                for for_export in (False, True):
                    with self.assertRaises(HTTPException) as error:
                        service.build_report(self.user, self.filters, for_export=for_export)
                    self.assertEqual(error.exception.status_code, 422)
        self.db.salesDb.documents.clear()
        self.sale([], precioTotal=0, precioTotalOriginal=0)
        del self.db.salesDb.documents[0]["productos"]
        with self.assertRaises(HTTPException):
            service.build_report(self.user, self.filters)

    def test_empty_zero_omissions_respect_local_lima_dates_state_and_annulment(self):
        for options in ({"instant": "2026-01-01T04:59:59+00:00"}, {"local": 2}, {"anulado": True},
                        {"instant": "2026-03-01T05:00:00+00:00"}):
            self.sale([], precioTotal=0, precioTotalOriginal=0, **options)
        self.sale([], instant="2026-01-01T05:00:00+00:00", precioTotal=0, precioTotalOriginal=0)
        self.sale([], instant="2026-02-01T04:59:59", estado="credito", precioTotal=0, precioTotalOriginal=0)
        self.sale([], instant="2026-03-01T04:59:59+00:00", precioTotal=0, precioTotalOriginal=0)
        report = service.build_report(self.user, self.filters)
        self.assertEqual(report["resumen"]["ventasVaciasOmitidas"], 3)
        self.assertEqual([row["ventasVaciasOmitidas"] for row in report["mensual"]], [2, 1])
        paid = service.build_report(self.user, service.ReportFilters(date(2026, 1, 1), date(2026, 2, 28), ventas="pagadas"))
        self.assertEqual(paid["resumen"]["ventasVaciasOmitidas"], 2)
        self.assertEqual([row["ventasVaciasOmitidas"] for row in paid["mensual"]], [1, 1])
        self.assertEqual(paid["resumen"]["numeroVentas"], 0)

    def test_empty_sales_are_not_attributed_to_product_or_category_filters(self):
        self.seed()
        self.sale([], precioTotal=0, precioTotalOriginal=0)
        for filters in (service.ReportFilters(date(2026, 1, 1), date(2026, 2, 28), producto="inexistente"),
                        service.ReportFilters(date(2026, 1, 1), date(2026, 2, 28), categoria="inexistente")):
            report = service.build_report(self.user, filters)
            self.assertEqual(report["resumen"]["ventasVaciasOmitidas"], 1)
            self.assertEqual(report["resumen"]["numeroVentas"], 0)
            self.assertIn("sin asignación de producto o categoría", report["avisos"][1])

    def test_nonempty_sales_with_zero_header_total_still_use_their_lines(self):
        self.sale(precioTotal=0, precioTotalOriginal=0)
        report = service.build_report(self.user, self.filters)
        self.assertEqual(report["resumen"]["ventasVaciasOmitidas"], 0)
        self.assertEqual(report["resumen"]["importeVendido"], 40)
        self.assertEqual(report["resumen"]["numeroVentas"], 1)

    async def test_empty_zero_sales_api_pagination_and_monthly_excel_notices(self):
        self.sale()
        self.sale([], precioTotal=0, precioTotalOriginal=0)
        self.sale([], precioTotal=0, precioTotalOriginal=0)
        self.sale([], instant="2026-02-14T18:33:36+00:00", precioTotal=0, precioTotalOriginal=0)
        for scope in ("todas", "pagadas"):
            for page in (1, 2):
                status, _, body = await self.request(ventas=scope, xpage=1, page=page, fecha_hasta="2026-03-31")
                self.assertEqual(status, 200)
                result = json.loads(body)["data"]
                self.assertEqual(result["resumen"]["ventasVaciasOmitidas"], 3)
                self.assertEqual(result["resumen"]["numeroVentas"], 1)
                self.assertEqual(result["resumen"]["importeVendido"], 40)
                self.assertIn("Ventas vacías omitidas: 3.", result["avisos"][1])
            status, _, body = await self.request("/api/reports/sales/excel", ventas=scope, fecha_hasta="2026-03-31")
            self.assertEqual(status, 200)
            book = load_workbook(BytesIO(body))
            self.assertEqual(book.sheetnames, ["enero2026", "febrero2026", "marzo2026"])
            january, february, march = book.worksheets
            self.assertEqual(january.tables["Ventas_2026_01"].ref, "A1:I2")
            self.assertEqual(january["I4"].value, 40)
            self.assertIn("Ventas vacías omitidas: 2.", january["A8"].value)
            self.assertIn("Ventas vacías omitidas: 1.", february["A8"].value)
            for sheet in (january, february):
                self.assertIn("$I$8", str(sheet.print_area))
                self.assertEqual(sheet["G1"].value, "P. COMPRA")
                self.assertEqual(sheet["H1"].value, "P. VENTA")
                self.assertEqual(sheet.freeze_panes, "A2")
            self.assertFalse(any("omitidas" in str(cell.value) for row in march for cell in row))
            self.assertIn("$I$7", str(march.print_area))

    async def test_only_empty_zero_sales_can_be_viewed_and_exported(self):
        self.sale([], precioTotal=0, precioTotalOriginal=0)
        status, _, body = await self.request()
        self.assertEqual(status, 200)
        report = json.loads(body)["data"]
        self.assertEqual(report["filas"], [])
        self.assertEqual(report["resumen"]["numeroVentas"], 0)
        self.assertEqual(report["resumen"]["ventasVaciasOmitidas"], 1)
        self.assertTrue(all(month["importeVendido"] == 0 for month in report["mensual"]))
        status, _, body = await self.request("/api/reports/sales/excel")
        self.assertEqual(status, 200)
        sheet = load_workbook(BytesIO(body))["enero2026"]
        self.assertIn("Sin ventas", sheet["B2"].value)
        self.assertEqual(sheet["I4"].value, 0)
        self.assertIn("Ventas vacías omitidas: 1.", sheet["A8"].value)

    async def test_empty_sale_with_original_amount_is_blocked_on_both_routes(self):
        sale = self.sale([], precioTotal=0, precioTotalOriginal=40)
        for path in ("/api/reports/sales", "/api/reports/sales/excel"):
            status, _, body = await self.request(path)
            self.assertEqual(status, 422)
            self.assertIn(str(sale["_id"]), json.loads(body)["detail"])

    def test_ranking_supports_quantity_and_amount(self):
        for index in range(12):
            self.sale([self.line(quantity=index + 1, price=1, product={"_id": ObjectId()}, name=f"Producto {index:02}")])
        self.sale([self.line(quantity=1, price=1000, name="Más importe")])
        report = service.build_report(self.user, self.filters)
        self.assertEqual(len(report["topCantidad"]), 10)
        self.assertEqual(report["topCantidad"][0]["cantidad"], 12)
        self.assertEqual(report["topImporte"][0]["importeVendido"], 1000)

    def test_period_limits_and_year_transition(self):
        with self.assertRaises(HTTPException):
            service.ReportFilters(date(2026, 2, 1), date(2026, 1, 1))
        with self.assertRaises(HTTPException):
            service.ReportFilters(date(2020, 1, 1), date(2026, 1, 1))
        report = service.build_report(self.user, service.ReportFilters(date(2025, 12, 1), date(2026, 1, 31)))
        self.assertEqual([row["mes"] for row in report["mensual"]], ["2025-12", "2026-01"])

    def test_safety_limit_does_not_truncate_silently(self):
        self.seed()
        with patch.object(service, "MAX_LINES", 1), self.assertRaises(HTTPException) as error:
            service.build_report(self.user, self.filters)
        self.assertEqual(error.exception.status_code, 422)

    def test_reporting_does_not_mutate_any_documents(self):
        self.seed()
        before = deepcopy((self.db.salesDb.documents, self.db.productsDb.documents))
        service.export_excel(service.build_report(self.user, self.filters, True))
        self.assertEqual(before, (self.db.salesDb.documents, self.db.productsDb.documents))

    def test_excel_memory_limit_is_explicit_and_never_exports_partial_rows(self):
        self.seed()
        with patch.object(service, "MAX_EXCEL_LINES", 1):
            self.assertEqual(service.build_report(self.user, self.filters)["totalFilas"], 2)
            with self.assertRaises(HTTPException) as error:
                service.build_report(self.user, self.filters, True)
            self.assertIn("Excel supera 1 líneas", error.exception.detail)

    async def test_pagination_does_not_change_totals_charts_or_export(self):
        self.seed()
        status, headers, body = await self.request(xpage=1, page=2)
        self.assertEqual(status, 200)
        self.assertEqual(headers[b"cache-control"], b"no-store")
        report = json.loads(body)["data"]
        self.assertEqual(len(report["filas"]), 1)
        self.assertEqual(report["totalFilas"], 2)
        self.assertEqual(report["resumen"]["importeVendido"], 112)
        self.assertEqual(report["mensual"][0]["importeVendido"], 112)
        self.assertNotIn("detalle", report)
        status, headers, body = await self.request("/api/reports/sales/excel")
        self.assertEqual(status, 200)
        self.assertIn(b"ventas_local_1_2026-01-01_2026-02-28.xlsx", headers[b"content-disposition"])
        book = load_workbook(BytesIO(body))
        self.assertEqual(book.sheetnames, ["enero2026", "febrero2026"])
        january, february = book.worksheets
        self.assertEqual(january.tables["Ventas_2026_01"].ref, "A1:I3")
        self.assertEqual(sum(january.cell(row, 9).value for row in (2, 3)), 112)
        self.assertEqual(january["I5"].value, 112)
        self.assertEqual(january["G3"].value, 15)
        self.assertEqual(january["H3"].value, 21.2)
        self.assertEqual(january["G2"].value, 0)
        self.assertEqual(january.freeze_panes, "A2")
        self.assertEqual(january["G3"].data_type, "n")
        self.assertEqual(february["I4"].value, 0)
        self.assertIn("Sin ventas", february["B2"].value)
        for sheet in book:
            self.assertFalse(any("referencial" in str(cell.value) for row in sheet for cell in row))

    async def test_excel_uses_same_product_filter_and_protects_formula_text(self):
        name = '=HYPERLINK("https://example.invalid","test")'
        self.sale([self.line(name=name), self.line(1, 6, self.other, "Bolsa")])
        status, _, body = await self.request("/api/reports/sales/excel", producto="HYPERLINK")
        self.assertEqual(status, 200)
        book = load_workbook(BytesIO(body), data_only=False)
        self.assertEqual(book.worksheets[0]["B2"].value, name)
        self.assertEqual(book.worksheets[0]["B2"].data_type, "s")
        self.assertEqual(book.worksheets[0]["I2"].value, 40)
        self.assertEqual(book.worksheets[0].tables["Ventas_2026_01"].ref, "A1:I2")

    async def test_missing_cost_excel_is_not_zero(self):
        self.sale([self.line(product={"_id": ObjectId()})])
        status, _, body = await self.request("/api/reports/sales/excel")
        self.assertEqual(status, 200)
        book = load_workbook(BytesIO(body))
        self.assertEqual(book.worksheets[0]["G2"].value, "No disponible")
        for column in "CDEF":
            self.assertIsNone(book.worksheets[0][f"{column}2"].value)

    async def test_authentication_permissions_and_local_isolation_on_both_routes(self):
        self.seed()
        for path in ("/api/reports/sales", "/api/reports/sales/excel"):
            self.assertEqual((await self.request(path, user=None))[0], 401)
            self.assertEqual((await self.request(path, token="invalid"))[0], 401)
            self.assertEqual((await self.request(path, user=self.reader))[0], 403)
            self.assertEqual((await self.request(path, local=2))[0], 403)
        status, _, body = await self.request(user=self.foreign)
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body)["data"]["resumen"]["numeroVentas"], 1)
        self.assertEqual(json.loads(body)["data"]["categorias"], [])

    async def test_invalid_queries_and_empty_excel(self):
        for params in ({"fecha_desde": "invalid"}, {"fecha_desde": "2026-03-01"}, {"xpage": 101},
                       {"page": 0}, {"producto": "x" * 121}, {"fecha_desde": "2020-01-01"}, {"ventas": "credito"}):
            self.assertEqual((await self.request(**params))[0], 422)
        status, _, body = await self.request("/api/reports/sales/excel")
        self.assertEqual(status, 200)
        book = load_workbook(BytesIO(body))
        for sheet in book:
            self.assertIn("Sin ventas", sheet["B2"].value)
            self.assertEqual(sheet["A4"].value, 0)
            self.assertEqual(sheet["I4"].value, 0)
            self.assertEqual(len(sheet.tables), 0)

    async def test_paid_filter_applies_to_summary_charts_ranking_pagination_and_excel(self):
        self.seed()
        self.sale(estado="pendiente")
        status, _, body = await self.request(ventas="pagadas", xpage=1)
        self.assertEqual(status, 200)
        report = json.loads(body)["data"]
        self.assertEqual(report["filtros"]["ventas"], "pagadas")
        self.assertEqual(report["resumen"]["importeVendido"], 46)
        self.assertEqual(report["resumen"]["unidadesVendidas"], 3)
        self.assertEqual(report["resumen"]["numeroVentas"], 1)
        self.assertEqual(report["mensual"][0]["importeVendido"], 46)
        self.assertEqual(sum(row["importeVendido"] for row in report["topImporte"]), 46)
        self.assertEqual(report["totalFilas"], 2)
        self.assertEqual(len(report["filas"]), 1)
        self.assertIn("completamente pagadas", report["avisos"][0])
        status, _, body = await self.request("/api/reports/sales/excel", ventas="pagadas")
        self.assertEqual(status, 200)
        january = load_workbook(BytesIO(body))["enero2026"]
        self.assertEqual(january.tables["Ventas_2026_01"].ref, "A1:I3")
        self.assertEqual(january["I5"].value, 46)
        self.assertIn("Solo pagadas", january["A7"].value)
        status, _, body = await self.request(ventas="todas")
        self.assertEqual(json.loads(body)["data"]["resumen"]["importeVendido"], 152)

    def test_fully_collected_credit_is_paid_by_sale_date_not_payment_date(self):
        self.seed()
        self.db.salesDb.documents[1].update(estado="cancelado", fechaCancelacion=datetime(2026, 8, 1))
        report = service.build_report(self.user, service.ReportFilters(date(2026, 1, 1), date(2026, 2, 28), ventas="pagadas"))
        self.assertEqual(report["resumen"]["importeVendido"], 112)
        self.assertEqual(report["mensual"][0]["numeroVentas"], 2)

    async def test_requested_excel_columns_and_direct_registered_purchase_price(self):
        self.db.productsDb.documents[0].update(nombre="20-20-20 CERES", unidadDeMedida="50 KILOGRAMOS",
            marca="", categoria="Fertilizantes", descripcion="", precioCompra=14.125678, precioVenta=999)
        self.sale([self.line(11, 158, name="20-20-20 CERES", precioBuy=777)])
        status, _, body = await self.request("/api/reports/sales/excel")
        self.assertEqual(status, 200)
        january = load_workbook(BytesIO(body))["enero2026"]
        self.assertEqual([cell.value for cell in january[1]], ["CANTIDAD", "PRODUCTO", "PRESENTACION", "MARCA",
            "CATEGORIA", "DESCRIPCION", "P. COMPRA", "P. VENTA", "IMPORTE VENDIDO"])
        self.assertEqual([cell.value for cell in january[2]], [11, "20-20-20 CERES", "50 KILOGRAMOS", None,
            "Fertilizantes", None, 14.125678, 158, 1738])
        self.assertEqual(january["G2"].data_type, "n")
        self.assertIn("promedio ponderado", january["H1"].comment.text)

    async def test_excel_sheets_follow_partial_date_range_across_years(self):
        for instant, quantity in (("2025-12-19T15:00:00+00:00", 99), ("2025-12-20T15:00:00+00:00", 1),
                                  ("2026-01-15T15:00:00+00:00", 2), ("2026-02-06T04:59:59+00:00", 3),
                                  ("2026-02-06T05:00:00+00:00", 99)):
            self.sale([self.line(quantity, 10)], instant=instant)
        status, _, body = await self.request("/api/reports/sales/excel", fecha_desde="2025-12-20", fecha_hasta="2026-02-05")
        self.assertEqual(status, 200)
        book = load_workbook(BytesIO(body))
        self.assertEqual(book.sheetnames, ["diciembre2025", "enero2026", "febrero2026"])
        self.assertEqual([sheet["A2"].value for sheet in book], [1, 2, 3])
        self.assertEqual([sheet["I4"].value for sheet in book], [10, 20, 30])
        self.assertEqual(sum(sheet["I4"].value for sheet in book), 60)

    async def test_same_month_in_different_years_has_distinct_sheets(self):
        self.sale(instant="2025-01-15T15:00:00+00:00")
        self.sale(instant="2026-01-15T15:00:00+00:00")
        status, _, body = await self.request("/api/reports/sales/excel", fecha_desde="2025-01-01", fecha_hasta="2026-01-31")
        self.assertEqual(status, 200)
        book = load_workbook(BytesIO(body))
        self.assertEqual(len(book.sheetnames), 13)
        self.assertEqual(book.sheetnames[0], "enero2025")
        self.assertEqual(book.sheetnames[-1], "enero2026")
        self.assertEqual(book["enero2025"]["I2"].value, 40)
        self.assertEqual(book["enero2026"]["I2"].value, 40)

    async def test_exported_product_metadata_is_literal_text(self):
        self.db.productsDb.documents[0].update(unidadDeMedida="=1+1", marca="=2+2", descripcion="=3+3")
        self.sale()
        status, _, body = await self.request("/api/reports/sales/excel")
        self.assertEqual(status, 200)
        january = load_workbook(BytesIO(body), data_only=False)["enero2026"]
        for column, value in (("C", "=1+1"), ("D", "=2+2"), ("F", "=3+3")):
            self.assertEqual(january[f"{column}2"].value, value)
            self.assertEqual(january[f"{column}2"].data_type, "s")


if __name__ == "__main__":
    unittest.main()
